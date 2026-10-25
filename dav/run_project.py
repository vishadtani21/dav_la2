"""
Student Academic Performance Analysis Using Python and Excel
End-to-End Execution Script
"""

import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.model_selection import train_test_split
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import joblib

# Set random seed for reproducibility
np.random.seed(42)

# Output directory
output_dir = r"c:\Users\Ajit Adtani\OneDrive\Desktop\dav"
os.makedirs(output_dir, exist_ok=True)
charts_dir = os.path.join(output_dir, "charts")
os.makedirs(charts_dir, exist_ok=True)

# -----------------------------------------------------------------------------
# STEP 1: SYNTHETIC DATASET CREATION (320 Records)
# -----------------------------------------------------------------------------
n_students = 320

student_ids = [f"STU{1001 + i}" for i in range(n_students)]
genders = np.random.choice(["Male", "Female"], size=n_students, p=[0.52, 0.48])

# Underlying academic characteristics with realistic correlations
attendance = np.random.normal(loc=76, scale=12, size=n_students)
attendance = np.clip(attendance, 40, 100).round(1)

study_hours = np.random.normal(loc=4.5, scale=1.8, size=n_students)
study_hours = np.clip(study_hours, 0.5, 12).round(1)

# Assignment score correlated with attendance and study hours
assignment_score = (
    0.35 * attendance + 3.2 * study_hours + np.random.normal(0, 7, size=n_students) + 30
)
assignment_score = np.clip(assignment_score, 35, 100).round(1)

# Internal marks correlated with assignment score and study hours
internal_marks = (
    0.45 * assignment_score + 0.25 * attendance + 1.5 * study_hours + np.random.normal(0, 6, size=n_students) + 10
)
internal_marks = np.clip(internal_marks, 30, 100).round(1)

# Previous semester marks
previous_marks = np.random.normal(loc=70, scale=12, size=n_students)
previous_marks = np.clip(previous_marks, 40, 98).round(1)

# Final Examination marks - dependent realistic relationship
final_marks = (
    0.22 * attendance +
    1.8 * study_hours +
    0.25 * assignment_score +
    0.30 * internal_marks +
    0.20 * previous_marks +
    np.random.normal(0, 4.5, size=n_students) - 8
)
final_marks = np.clip(final_marks, 25, 100).round(1)

df_clean_base = pd.DataFrame({
    "Student_ID": student_ids,
    "Gender": genders,
    "Attendance_Percentage": attendance,
    "Study_Hours": study_hours,
    "Assignment_Score": assignment_score,
    "Internal_Marks": internal_marks,
    "Previous_Semester_Marks": previous_marks,
    "Final_Marks": final_marks
})

# Create Raw Dataset with deliberate issues (Duplicates, Missing Values, Invalid Out-of-range Values)
df_raw = df_clean_base.copy()

# 1. Add 6 duplicate rows
duplicates = df_raw.iloc[[12, 45, 88, 120, 175, 230]].copy()
df_raw = pd.concat([df_raw, duplicates], ignore_index=True)

# 2. Add missing values (NaN) in a few columns
np.random.seed(99)
missing_att_idx = [15, 62, 140, 210]
missing_study_idx = [33, 95, 180]
missing_assign_idx = [50, 125, 270]
missing_prev_idx = [72, 199]

for idx in missing_att_idx:
    df_raw.loc[idx, "Attendance_Percentage"] = np.nan
for idx in missing_study_idx:
    df_raw.loc[idx, "Study_Hours"] = np.nan
for idx in missing_assign_idx:
    df_raw.loc[idx, "Assignment_Score"] = np.nan
for idx in missing_prev_idx:
    df_raw.loc[idx, "Previous_Semester_Marks"] = np.nan

# 3. Add deliberate invalid out-of-range values to demonstrate cleaning
df_raw.loc[25, "Attendance_Percentage"] = 145.0  # Invalid attendance > 100%
df_raw.loc[82, "Study_Hours"] = -3.5             # Invalid negative study hours
df_raw.loc[160, "Assignment_Score"] = 125.0       # Invalid score > 100
df_raw.loc[215, "Internal_Marks"] = -10.0         # Invalid negative marks

# Save raw dataset
raw_csv_path = os.path.join(output_dir, "raw_student_data.csv")
df_raw.to_csv(raw_csv_path, index=False)
print(f"Raw dataset created with {len(df_raw)} records and saved to {raw_csv_path}")

# -----------------------------------------------------------------------------
# STEP 2: DATA CLEANING & PREPROCESSING
# -----------------------------------------------------------------------------
df_clean = df_raw.copy()

# A. Duplicate Removal
duplicates_count = df_clean.duplicated().sum()
df_clean = df_clean.drop_duplicates().reset_index(drop=True)
print(f"Removed {duplicates_count} duplicate records.")

# B. Invalid Values Handling
# Attendance must be 0-100
invalid_att = (df_clean["Attendance_Percentage"] < 0) | (df_clean["Attendance_Percentage"] > 100)
df_clean.loc[invalid_att, "Attendance_Percentage"] = np.nan

# Study hours must be >= 0 and <= 16
invalid_study = (df_clean["Study_Hours"] < 0) | (df_clean["Study_Hours"] > 16)
df_clean.loc[invalid_study, "Study_Hours"] = np.nan

# Assignment score must be 0-100
invalid_assign = (df_clean["Assignment_Score"] < 0) | (df_clean["Assignment_Score"] > 100)
df_clean.loc[invalid_assign, "Assignment_Score"] = np.nan

# Internal marks must be 0-100
invalid_internal = (df_clean["Internal_Marks"] < 0) | (df_clean["Internal_Marks"] > 100)
df_clean.loc[invalid_internal, "Internal_Marks"] = np.nan

# Previous semester marks must be 0-100
invalid_prev = (df_clean["Previous_Semester_Marks"] < 0) | (df_clean["Previous_Semester_Marks"] > 100)
df_clean.loc[invalid_prev, "Previous_Semester_Marks"] = np.nan

# C. Missing Value Imputation using Median (robust against skew)
numeric_cols = ["Attendance_Percentage", "Study_Hours", "Assignment_Score", "Internal_Marks", "Previous_Semester_Marks", "Final_Marks"]

missing_summary_before = df_clean[numeric_cols].isnull().sum()
print("Missing values before imputation:")
print(missing_summary_before)

for col in numeric_cols:
    if df_clean[col].isnull().sum() > 0:
        median_val = df_clean[col].median()
        df_clean[col] = df_clean[col].fillna(median_val)

# D. Data Type Casting & Rounding
df_clean["Student_ID"] = df_clean["Student_ID"].astype(str)
df_clean["Gender"] = df_clean["Gender"].astype(str)
for col in numeric_cols:
    df_clean[col] = df_clean[col].astype(float).round(2)

# E. Outlier Detection using IQR Method
outlier_report = {}
for col in numeric_cols:
    q1 = df_clean[col].quantile(0.25)
    q3 = df_clean[col].quantile(0.75)
    iqr = q3 - q1
    lower_bound = q1 - 1.5 * iqr
    upper_bound = q3 + 1.5 * iqr
    outliers = df_clean[(df_clean[col] < lower_bound) | (df_clean[col] > upper_bound)]
    outlier_report[col] = {
        "Q1": round(q1, 2),
        "Q3": round(q3, 2),
        "IQR": round(iqr, 2),
        "Lower_Bound": round(lower_bound, 2),
        "Upper_Bound": round(upper_bound, 2),
        "Outlier_Count": len(outliers)
    }

# Save cleaned CSV
cleaned_csv_path = os.path.join(output_dir, "cleaned_student_data.csv")
df_clean.to_csv(cleaned_csv_path, index=False)
print(f"Cleaned dataset saved with {len(df_clean)} records to {cleaned_csv_path}")

# -----------------------------------------------------------------------------
# STEP 3: STATISTICAL ANALYSIS & PERFORMANCE CATEGORIES
# -----------------------------------------------------------------------------
stats_list = []
for col in numeric_cols:
    s = df_clean[col]
    mean_v = s.mean()
    median_v = s.median()
    mode_v = s.mode().iloc[0]
    std_v = s.std()
    var_v = s.var()
    min_v = s.min()
    max_v = s.max()
    range_v = max_v - min_v
    q1_v = s.quantile(0.25)
    q3_v = s.quantile(0.75)
    
    stats_list.append({
        "Variable": col,
        "Mean": round(mean_v, 2),
        "Median": round(median_v, 2),
        "Mode": round(mode_v, 2),
        "Std_Dev": round(std_v, 2),
        "Variance": round(var_v, 2),
        "Min": round(min_v, 2),
        "Max": round(max_v, 2),
        "Range": round(range_v, 2),
        "Q1 (25%)": round(q1_v, 2),
        "Q3 (75%)": round(q3_v, 2)
    })

df_stats = pd.DataFrame(stats_list)
print("\nDescriptive Statistical Summary:")
print(df_stats.to_string(index=False))

# Performance Categorization
def assign_category(score):
    if score >= 85.0:
        return "Excellent (85-100)"
    elif score >= 70.0:
        return "Good (70-84.99)"
    elif score >= 50.0:
        return "Average (50-69.99)"
    else:
        return "Needs Improvement (<50)"

df_clean["Performance_Category"] = df_clean["Final_Marks"].apply(assign_category)

category_order = ["Excellent (85-100)", "Good (70-84.99)", "Average (50-69.99)", "Needs Improvement (<50)"]
category_counts = df_clean["Performance_Category"].value_counts().reindex(category_order, fill_value=0)
category_pcts = ((category_counts / len(df_clean)) * 100).round(2)
category_avg_marks = df_clean.groupby("Performance_Category")["Final_Marks"].mean().reindex(category_order, fill_value=0).round(2)

df_categories = pd.DataFrame({
    "Performance_Category": category_order,
    "Student_Count": category_counts.values,
    "Percentage": category_pcts.values,
    "Average_Final_Marks": category_avg_marks.values
})
print("\nPerformance Categories Distribution:")
print(df_categories.to_string(index=False))

# -----------------------------------------------------------------------------
# STEP 4: CORRELATION ANALYSIS
# -----------------------------------------------------------------------------
df_corr = df_clean[numeric_cols].corr().round(4)
print("\nCorrelation Matrix:")
print(df_corr)

corr_with_final = df_corr["Final_Marks"].drop("Final_Marks").sort_values(ascending=False)
print("\nCorrelation with Final Marks:")
print(corr_with_final)

# -----------------------------------------------------------------------------
# STEP 5: DATA VISUALIZATION (7 CHARTS)
# -----------------------------------------------------------------------------
plt.style.use('seaborn-v0_8-whitegrid' if 'seaborn-v0_8-whitegrid' in plt.style.available else 'default')
palette = ["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728", "#9467bd"]

# Chart 1: Histogram of Final Marks
fig1, ax1 = plt.subplots(figsize=(8, 5), dpi=300)
sns.histplot(df_clean["Final_Marks"], kde=True, bins=15, color="#1e3d59", edgecolor="white", ax=ax1)
ax1.axvline(df_clean["Final_Marks"].mean(), color="#ff6e40", linestyle="--", linewidth=2, label=f"Mean: {df_clean['Final_Marks'].mean():.2f}")
ax1.axvline(df_clean["Final_Marks"].median(), color="#17b978", linestyle="-.", linewidth=2, label=f"Median: {df_clean['Final_Marks'].median():.2f}")
ax1.set_title("1. Distribution of Final Examination Marks", fontsize=14, fontweight="bold", pad=12)
ax1.set_xlabel("Final Examination Marks", fontsize=12)
ax1.set_ylabel("Student Frequency", fontsize=12)
ax1.legend(loc="upper right", frameon=True)
plt.tight_layout()
chart1_path = os.path.join(charts_dir, "01_final_marks_distribution.png")
fig1.savefig(chart1_path)
plt.close(fig1)

# Chart 2: Bar Chart of Performance Categories
fig2, ax2 = plt.subplots(figsize=(8, 5), dpi=300)
bars = ax2.bar(df_categories["Performance_Category"], df_categories["Average_Final_Marks"], color=["#2e7d32", "#1976d2", "#f57c00", "#c62828"], edgecolor="black", alpha=0.85)
for bar in bars:
    yval = bar.get_height()
    ax2.text(bar.get_x() + bar.get_width()/2.0, yval + 1.2, f"{yval:.1f}", ha='center', va='bottom', fontsize=11, fontweight="bold")
ax2.set_title("2. Average Marks by Student Performance Category", fontsize=14, fontweight="bold", pad=12)
ax2.set_xlabel("Performance Category", fontsize=12)
ax2.set_ylabel("Average Final Marks", fontsize=12)
ax2.set_ylim(0, 105)
plt.xticks(rotation=15, ha='right')
plt.tight_layout()
chart2_path = os.path.join(charts_dir, "02_performance_categories_barchart.png")
fig2.savefig(chart2_path)
plt.close(fig2)

# Chart 3: Scatter Plot - Study Hours vs Final Marks
fig3, ax3 = plt.subplots(figsize=(8, 5), dpi=300)
sns.regplot(data=df_clean, x="Study_Hours", y="Final_Marks", scatter_kws={'alpha':0.6, 'color': '#0d47a1'}, line_kws={'color': '#d50000', 'linewidth': 2}, ax=ax3)
ax3.set_title("3. Relationship: Daily Study Hours vs. Final Marks", fontsize=14, fontweight="bold", pad=12)
ax3.set_xlabel("Daily Study Hours", fontsize=12)
ax3.set_ylabel("Final Examination Marks", fontsize=12)
plt.tight_layout()
chart3_path = os.path.join(charts_dir, "03_study_hours_vs_final_marks.png")
fig3.savefig(chart3_path)
plt.close(fig3)

# Chart 4: Scatter Plot - Attendance vs Final Marks
fig4, ax4 = plt.subplots(figsize=(8, 5), dpi=300)
sns.regplot(data=df_clean, x="Attendance_Percentage", y="Final_Marks", scatter_kws={'alpha':0.6, 'color': '#00695c'}, line_kws={'color': '#ff6d00', 'linewidth': 2}, ax=ax4)
ax4.set_title("4. Relationship: Attendance Percentage vs. Final Marks", fontsize=14, fontweight="bold", pad=12)
ax4.set_xlabel("Attendance Percentage (%)", fontsize=12)
ax4.set_ylabel("Final Examination Marks", fontsize=12)
plt.tight_layout()
chart4_path = os.path.join(charts_dir, "04_attendance_vs_final_marks.png")
fig4.savefig(chart4_path)
plt.close(fig4)

# Chart 5: Correlation Heatmap
fig5, ax5 = plt.subplots(figsize=(8, 6), dpi=300)
mask = np.triu(np.ones_like(df_corr, dtype=bool), k=1)
sns.heatmap(df_corr, annot=True, fmt=".2f", cmap="vlag", center=0, square=True, linewidths=0.5, cbar_kws={"shrink": 0.8}, ax=ax5)
ax5.set_title("5. Academic Variables Correlation Heatmap", fontsize=14, fontweight="bold", pad=12)
plt.xticks(rotation=30, ha='right')
plt.yticks(rotation=0)
plt.tight_layout()
chart5_path = os.path.join(charts_dir, "05_correlation_heatmap.png")
fig5.savefig(chart5_path)
plt.close(fig5)

# Chart 6: Box Plot of Academic Variables
fig6, ax6 = plt.subplots(figsize=(9, 5), dpi=300)
box_data = df_clean[["Attendance_Percentage", "Assignment_Score", "Internal_Marks", "Previous_Semester_Marks", "Final_Marks"]]
sns.boxplot(data=box_data, palette="Set2", ax=ax6)
ax6.set_title("6. Box Plot of Examination and Academic Variables", fontsize=14, fontweight="bold", pad=12)
ax6.set_ylabel("Marks / Score Range", fontsize=12)
plt.xticks(rotation=20, ha='right')
plt.tight_layout()
chart6_path = os.path.join(charts_dir, "06_academic_variables_boxplot.png")
fig6.savefig(chart6_path)
plt.close(fig6)

# -----------------------------------------------------------------------------
# STEP 6: LINEAR REGRESSION MODELING
# -----------------------------------------------------------------------------
feature_cols = [
    "Attendance_Percentage",
    "Study_Hours",
    "Assignment_Score",
    "Internal_Marks",
    "Previous_Semester_Marks"
]
target_col = "Final_Marks"

X = df_clean[feature_cols]
y = df_clean[target_col]

# 80-20 Train-Test Split with reproducible seed
X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.20, random_state=42)

model = LinearRegression()
model.fit(X_train, y_train)

# Predictions
y_pred_train = model.predict(X_train)
y_pred_test = model.predict(X_test)

# Metrics
mae = mean_absolute_error(y_test, y_pred_test)
mse = mean_squared_error(y_test, y_pred_test)
rmse = np.sqrt(mse)
r2 = r2_score(y_test, y_pred_test)
train_r2 = r2_score(y_train, y_pred_train)

df_metrics = pd.DataFrame({
    "Evaluation_Metric": [
        "Mean Absolute Error (MAE)",
        "Mean Squared Error (MSE)",
        "Root Mean Squared Error (RMSE)",
        "R-squared Score (R2 Test)",
        "R-squared Score (R2 Train)"
    ],
    "Value": [round(mae, 4), round(mse, 4), round(rmse, 4), round(r2, 4), round(train_r2, 4)],
    "Interpretation": [
        "Average absolute difference between actual & predicted marks",
        "Average squared difference, penalizes larger errors",
        "Standard deviation of unexplained variance in marks",
        f"Model explains {r2*100:.2f}% of final marks variance on unseen test data",
        f"Model explains {train_r2*100:.2f}% of final marks variance on training data"
    ]
})

print("\nModel Evaluation Metrics:")
print(df_metrics.to_string(index=False))

# Coefficients
coefficients = pd.DataFrame({
    "Feature": ["Intercept (Constant)"] + feature_cols,
    "Coefficient": [round(model.intercept_, 4)] + [round(c, 4) for c in model.coef_],
    "Impact_Description": [
        "Baseline score when all independent features are zero",
        "Marks gain per 1% increase in attendance",
        "Marks gain per 1 additional daily study hour",
        "Marks gain per 1 point increase in assignment score",
        "Marks gain per 1 point increase in internal marks",
        "Marks gain per 1 point increase in previous semester marks"
    ]
})
print("\nRegression Equation Coefficients:")
print(coefficients.to_string(index=False))

# Test Predictions Comparison Table
df_predictions = X_test.copy()
df_predictions["Actual_Final_Marks"] = y_test.values
df_predictions["Predicted_Final_Marks"] = np.round(y_pred_test, 2)
df_predictions["Prediction_Error"] = np.round(df_predictions["Actual_Final_Marks"] - df_predictions["Predicted_Final_Marks"], 2)
df_predictions["Absolute_Error"] = np.abs(df_predictions["Prediction_Error"])

# Chart 7: Actual vs Predicted Marks Plot
fig7, ax7 = plt.subplots(figsize=(8, 6), dpi=300)
ax7.scatter(y_test, y_pred_test, color="#1565c0", alpha=0.7, edgecolors="k", s=50, label="Test Students")
# Ideal 45 degree line
min_val = min(y_test.min(), y_pred_test.min()) - 2
max_val = max(y_test.max(), y_pred_test.max()) + 2
ax7.plot([min_val, max_val], [min_val, max_val], color="#c62828", linestyle="--", linewidth=2, label="Ideal Fit (Actual = Predicted)")
ax7.set_title(f"7. Actual vs. Predicted Final Marks (R² = {r2:.3f}, RMSE = {rmse:.2f})", fontsize=14, fontweight="bold", pad=12)
ax7.set_xlabel("Actual Final Examination Marks", fontsize=12)
ax7.set_ylabel("Predicted Final Examination Marks", fontsize=12)
ax7.set_xlim(min_val, max_val)
ax7.set_ylim(min_val, max_val)
ax7.legend(loc="upper left", frameon=True)
plt.tight_layout()
chart7_path = os.path.join(charts_dir, "07_actual_vs_predicted_marks.png")
fig7.savefig(chart7_path)
plt.close(fig7)

# Save model
model_path = os.path.join(output_dir, "student_performance_model.pkl")
joblib.dump(model, model_path)
print(f"Trained model saved to {model_path}")

# -----------------------------------------------------------------------------
# STEP 7: NEW STUDENT PREDICTION FUNCTION
# -----------------------------------------------------------------------------
def predict_new_student(attendance_pct, daily_study_hrs, assignment_sc, internal_mk, prev_sem_mk):
    """
    Validates input parameters and computes estimated final examination marks.
    """
    # Validation
    if not (0 <= attendance_pct <= 100):
        raise ValueError(f"Invalid Attendance: {attendance_pct}%. Must be between 0 and 100.")
    if not (0 <= daily_study_hrs <= 16):
        raise ValueError(f"Invalid Daily Study Hours: {daily_study_hrs}. Must be between 0 and 16.")
    if not (0 <= assignment_sc <= 100):
        raise ValueError(f"Invalid Assignment Score: {assignment_sc}. Must be between 0 and 100.")
    if not (0 <= internal_mk <= 100):
        raise ValueError(f"Invalid Internal Marks: {internal_mk}. Must be between 0 and 100.")
    if not (0 <= prev_sem_mk <= 100):
        raise ValueError(f"Invalid Previous Semester Marks: {prev_sem_mk}. Must be between 0 and 100.")
    
    input_data = pd.DataFrame([{
        "Attendance_Percentage": attendance_pct,
        "Study_Hours": daily_study_hrs,
        "Assignment_Score": assignment_sc,
        "Internal_Marks": internal_mk,
        "Previous_Semester_Marks": prev_sem_mk
    }])
    
    raw_pred = model.predict(input_data)[0]
    final_pred = round(float(np.clip(raw_pred, 0, 100)), 2)
    return final_pred

sample_pred = predict_new_student(85.0, 5.0, 88.0, 82.0, 78.0)
print(f"\nSample Prediction for Student (Att=85%, Study=5hrs, Assign=88, Int=82, Prev=78): {sample_pred}/100")

# -----------------------------------------------------------------------------
# STEP 8: EXCEL WORKBOOK GENERATION (OpenPyXL)
# -----------------------------------------------------------------------------
wb_path = os.path.join(output_dir, "Student_Academic_Performance_Analysis.xlsx")
wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Styling definitions
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
header_fill_teal = PatternFill(start_color="005A70", end_color="005A70", fill_type="solid")
title_font = Font(name="Calibri", size=14, bold=True, color="1B365D")
regular_font = Font(name="Calibri", size=11)
bold_font = Font(name="Calibri", size=11, bold=True)
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
double_bottom_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='double', color='1B365D')
)

def style_worksheet(ws, title_text, fill_color=header_fill_navy):
    # Sheet title
    ws.insert_rows(1, 2)
    ws["A1"] = title_text
    ws["A1"].font = title_font
    
    # Header row is now row 3
    header_row = 3
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = fill_color
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    # Data rows styling
    for r in range(4, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = regular_font
            cell.border = thin_border
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)
        
    ws.freeze_panes = "A4"

# 1. Sheet: Raw_Data
ws_raw = wb.create_sheet(title="Raw_Data")
for r_idx, row in enumerate([df_raw.columns.tolist()] + df_raw.values.tolist(), 1):
    ws_raw.append(row)
style_worksheet(ws_raw, "Raw Student Academic Dataset (Contains Deliberate Duplicates, Missing & Invalid Values)")

# 2. Sheet: Cleaned_Data
ws_clean = wb.create_sheet(title="Cleaned_Data")
for r_idx, row in enumerate([df_clean.columns.tolist()] + df_clean.values.tolist(), 1):
    ws_clean.append(row)
style_worksheet(ws_clean, "Cleaned Student Academic Dataset (320 Cleaned Records, Imputed & Verified)", header_fill_teal)

# Add Excel Formulas at bottom of Cleaned Data
max_r = ws_clean.max_row
summary_row = max_r + 2
ws_clean.cell(row=summary_row, column=2, value="AVERAGE (Excel Formula)").font = bold_font
ws_clean.cell(row=summary_row + 1, column=2, value="COUNT (Excel Formula)").font = bold_font

for col_idx, col_name in enumerate(df_clean.columns, 1):
    col_let = get_column_letter(col_idx)
    if col_name in numeric_cols:
        avg_cell = ws_clean.cell(row=summary_row, column=col_idx)
        avg_cell.value = f"=AVERAGE({col_let}4:{col_let}{max_r})"
        avg_cell.font = bold_font
        avg_cell.number_format = "0.00"
        avg_cell.border = double_bottom_border
        
        count_cell = ws_clean.cell(row=summary_row + 1, column=col_idx)
        count_cell.value = f"=COUNT({col_let}4:{col_let}{max_r})"
        count_cell.font = bold_font
        count_cell.border = double_bottom_border

# 3. Sheet: Statistical_Summary
ws_stats = wb.create_sheet(title="Statistical_Summary")
for r_idx, row in enumerate([df_stats.columns.tolist()] + df_stats.values.tolist(), 1):
    ws_stats.append(row)
style_worksheet(ws_stats, "Descriptive Statistics for All Academic Variables")

# 4. Sheet: Correlation_Matrix
ws_corr = wb.create_sheet(title="Correlation_Matrix")
corr_matrix_export = df_corr.reset_index().rename(columns={"index": "Variable"})
for r_idx, row in enumerate([corr_matrix_export.columns.tolist()] + corr_matrix_export.values.tolist(), 1):
    ws_corr.append(row)
style_worksheet(ws_corr, "Pearson Correlation Matrix of Academic Features", header_fill_teal)

# 5. Sheet: Performance_Categories
ws_cat = wb.create_sheet(title="Performance_Categories")
for r_idx, row in enumerate([df_categories.columns.tolist()] + df_categories.values.tolist(), 1):
    ws_cat.append(row)
style_worksheet(ws_cat, "Student Academic Performance Classification & Distribution")

# Add Total Row in Performance Categories
cat_max_r = ws_cat.max_row
total_r = cat_max_r + 1
ws_cat.cell(row=total_r, column=1, value="Total").font = bold_font
ws_cat.cell(row=total_r, column=2, value=f"=SUM(B4:B{cat_max_r})").font = bold_font
ws_cat.cell(row=total_r, column=3, value=f"=SUM(C4:C{cat_max_r})").font = bold_font
ws_cat.cell(row=total_r, column=4, value=f"=AVERAGE(D4:D{cat_max_r})").font = bold_font

# 6. Sheet: Model_Evaluation
ws_eval = wb.create_sheet(title="Model_Evaluation")
ws_eval.append(["Regression Equation & Parameters", "", ""])
ws_eval.append(["Feature Name", "Estimated Coefficient / Intercept", "Practical Interpretation"])
for _, row in coefficients.iterrows():
    ws_eval.append(row.tolist())
    
ws_eval.append([])
ws_eval.append(["Model Goodness-of-Fit Evaluation Metrics", "", ""])
ws_eval.append(["Metric Name", "Metric Score", "Metric Meaning & Range"])
for _, row in df_metrics.iterrows():
    ws_eval.append(row.tolist())
style_worksheet(ws_eval, "Linear Regression Model Specification & Evaluation Metrics", header_fill_navy)

# 7. Sheet: Predictions
ws_pred = wb.create_sheet(title="Predictions")
pred_export = df_predictions.reset_index().rename(columns={"index": "Original_Index"})
for r_idx, row in enumerate([pred_export.columns.tolist()] + pred_export.values.tolist(), 1):
    ws_pred.append(row)
style_worksheet(ws_pred, "Test Set Predictions vs Actual Final Examination Marks (64 Test Samples)", header_fill_teal)

# Save workbook
wb.save(wb_path)
print(f"Excel workbook generated successfully at: {wb_path}")

# Print automatic summary
print("="*75)
print("             AUTOMATIC MINI-PROJECT RESULTS SUMMARY")
print("="*75)
print(f"Total Raw Records Loaded       : {len(df_raw)}")
print(f"Duplicates Detected & Removed  : {duplicates_count}")
print(f"Total Cleaned Records          : {len(df_clean)}")
print(f"Average Final Examination Marks: {df_clean['Final_Marks'].mean():.2f}")
print(f"Average Attendance Percentage  : {df_clean['Attendance_Percentage'].mean():.2f}%")
print(f"Average Daily Study Hours      : {df_clean['Study_Hours'].mean():.2f} hrs/day")
print("\nPerformance Categories:")
for _, r in df_categories.iterrows():
    print(f" - {r['Performance_Category']:<25}: {r['Student_Count']:>3} students ({r['Percentage']:>5.1f}%) | Avg Marks: {r['Average_Final_Marks']}")
print("\nStrongest Correlations with Final Marks:")
for var, val in corr_with_final.items():
    print(f" - {var:<25}: r = {val:>6.4f}")
print("\nRegression Model Performance (Test Split = 20%):")
print(f" - Mean Absolute Error (MAE)   : {mae:.4f}")
print(f" - Mean Squared Error (MSE)    : {mse:.4f}")
print(f" - Root Mean Squared Error     : {rmse:.4f}")
print(f" - R-squared Score (R²)        : {r2:.4f} ({r2*100:.2f}% variance explained)")
print("="*75)
